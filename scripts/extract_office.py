#!/usr/bin/env python3
"""Extract DOCX and PPTX content to deterministic, searchable Markdown."""

from __future__ import annotations

import argparse
from collections.abc import Iterable
from pathlib import Path
import re
import sys


SUPPORTED_EXTENSIONS = {".docx", ".pptx"}
VISUAL_WARNING = (
    "[Visual or unsupported content detected; inspect the original file for "
    "images, charts, diagrams, equations, media, or spatial layout.]"
)
EMPTY_WARNING = (
    "[No extractable text; inspect the original file for visual or unsupported content.]"
)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Extract searchable DOCX or PPTX content to Markdown."
    )
    parser.add_argument("input", type=Path, help="DOCX or PPTX file to extract")
    parser.add_argument(
        "--output", required=True, type=Path, help="UTF-8 Markdown output path"
    )
    return parser.parse_args()


def normalize_text(text: str) -> str:
    """Normalize line endings and trailing whitespace without interpreting text."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in text.split("\n")]
    while lines and not lines[0]:
        lines.pop(0)
    while lines and not lines[-1]:
        lines.pop()
    return "\n".join(lines)


def markdown_cell(text: str) -> str:
    """Escape text for a Markdown table cell."""
    return (
        normalize_text(text)
        .replace("\\", "\\\\")
        .replace("|", "\\|")
        .replace("\n", "<br>")
    )


def markdown_table(rows: Iterable[Iterable[str]]) -> list[str]:
    matrix = [[markdown_cell(cell) for cell in row] for row in rows]
    if not matrix:
        return []
    width = max(len(row) for row in matrix)
    if width == 0:
        return []
    normalized = [row + [""] * (width - len(row)) for row in matrix]
    output = [
        "| " + " | ".join(normalized[0]) + " |",
        "| " + " | ".join(["---"] * width) + " |",
    ]
    output.extend("| " + " | ".join(row) + " |" for row in normalized[1:])
    return output


def extract_docx(input_path: Path) -> tuple[list[str], int]:
    try:
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:
        raise RuntimeError(
            "missing dependency 'python-docx'; run "
            "'python -m pip install -r requirements.txt'"
        ) from exc

    try:
        document = Document(input_path)
        chunks: list[str] = []
        extracted_blocks = 0
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph):
                text = normalize_text(block.text)
                if not text:
                    continue
                style_name = block.style.name if block.style is not None else ""
                heading = re.fullmatch(r"Heading ([1-9])", style_name)
                if heading:
                    level = min(int(heading.group(1)) + 1, 6)
                    chunks.extend([f"{'#' * level} {text}", ""])
                else:
                    chunks.extend([text, ""])
                extracted_blocks += 1
            elif isinstance(block, Table):
                table = markdown_table(
                    ([cell.text for cell in row.cells] for row in block.rows)
                )
                if table:
                    chunks.extend(table + [""])
                    extracted_blocks += 1

        warnings = 0
        visual_elements = {"drawing", "object", "oMath", "oMathPara", "pict"}
        has_visual_content = any(
            element.tag.rsplit("}", 1)[-1] in visual_elements
            for element in document.element.iter()
        )
        if has_visual_content:
            chunks.extend([VISUAL_WARNING, ""])
            warnings += 1
        if extracted_blocks == 0:
            chunks.extend([EMPTY_WARNING, ""])
            warnings += 1
        return chunks, warnings
    except Exception as exc:
        raise RuntimeError(f"could not extract DOCX: {exc}") from exc


def _pptx_shape_content(shape: object) -> tuple[list[str], bool]:
    """Return text/table Markdown and whether a shape needs visual inspection."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
        chunks: list[str] = []
        visual = False
        for child in shape.shapes:
            child_chunks, child_visual = _pptx_shape_content(child)
            chunks.extend(child_chunks)
            visual = visual or child_visual
        return chunks, visual

    if getattr(shape, "has_table", False):
        table = markdown_table(
            ([cell.text for cell in row.cells] for row in shape.table.rows)
        )
        return (table + [""] if table else []), False

    visual_xml_elements = {
        "audioFile",
        "blip",
        "chart",
        "graphicData",
        "oleObj",
        "videoFile",
    }
    has_visual_content = any(
        element.tag.rsplit("}", 1)[-1] in visual_xml_elements
        for element in shape.element.iter()
    )

    if getattr(shape, "has_text_frame", False):
        paragraphs = [
            normalize_text(paragraph.text)
            for paragraph in shape.text_frame.paragraphs
            if normalize_text(paragraph.text)
        ]
        return ([*paragraphs, ""] if paragraphs else []), has_visual_content

    visual_types = {
        MSO_SHAPE_TYPE.CHART,
        MSO_SHAPE_TYPE.DIAGRAM,
        MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT,
        MSO_SHAPE_TYPE.IGX_GRAPHIC,
        MSO_SHAPE_TYPE.INK,
        MSO_SHAPE_TYPE.LINKED_PICTURE,
        MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
        MSO_SHAPE_TYPE.MEDIA,
        MSO_SHAPE_TYPE.PICTURE,
        MSO_SHAPE_TYPE.WEB_VIDEO,
    }
    return [], has_visual_content or getattr(shape, "shape_type", None) in visual_types


def extract_pptx(input_path: Path) -> tuple[list[str], int]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise RuntimeError(
            "missing dependency 'python-pptx'; run "
            "'python -m pip install -r requirements.txt'"
        ) from exc

    try:
        presentation = Presentation(input_path)
        chunks: list[str] = []
        warnings = 0
        for slide_number, slide in enumerate(presentation.slides, start=1):
            chunks.extend([f"## Slide {slide_number}", ""])
            slide_chunks: list[str] = []
            visual = False
            for shape in slide.shapes:
                shape_chunks, shape_visual = _pptx_shape_content(shape)
                slide_chunks.extend(shape_chunks)
                visual = visual or shape_visual

            chunks.extend(slide_chunks)
            notes_extracted = False
            if slide.has_notes_slide:
                notes_frame = slide.notes_slide.notes_text_frame
                if notes_frame is not None:
                    notes = [
                        normalize_text(paragraph.text)
                        for paragraph in notes_frame.paragraphs
                        if normalize_text(paragraph.text)
                    ]
                    if notes:
                        chunks.extend(["### Speaker notes", "", *notes, ""])
                        notes_extracted = True

            if visual:
                chunks.extend([VISUAL_WARNING, ""])
                warnings += 1
            if not slide_chunks and not notes_extracted and not visual:
                chunks.extend([EMPTY_WARNING, ""])
                warnings += 1
        if not presentation.slides:
            chunks.extend([EMPTY_WARNING, ""])
            warnings += 1
        return chunks, warnings
    except Exception as exc:
        raise RuntimeError(f"could not extract PPTX: {exc}") from exc


def extract(input_path: Path, output_path: Path) -> int:
    if not input_path.is_file():
        raise ValueError(f"input file does not exist: {input_path}")
    extension = input_path.suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"input must be one of: {supported}")
    if input_path.resolve() == output_path.resolve():
        raise ValueError("input and output paths must be different")

    if extension == ".docx":
        body, warnings = extract_docx(input_path)
    else:
        body, warnings = extract_pptx(input_path)

    chunks = [f"# Extracted source: {output_path.stem}", "", *body]
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(
        "\n".join(chunks).rstrip() + "\n", encoding="utf-8", newline="\n"
    )
    return warnings


def main() -> int:
    args = parse_args()
    try:
        warnings = extract(args.input, args.output)
    except (OSError, RuntimeError, ValueError) as exc:
        print(f"error: {exc}", file=sys.stderr)
        return 1

    print(f"Extracted {args.input} to {args.output}.")
    if warnings:
        print(
            f"Warning: {warnings} location(s) need inspection in the original file.",
            file=sys.stderr,
        )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
