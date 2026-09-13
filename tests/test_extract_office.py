from __future__ import annotations

from pathlib import Path
import sys
import tempfile
import unittest


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))

from extract_office import VISUAL_WARNING, extract  # noqa: E402


class ExtractOfficeTests(unittest.TestCase):
    def test_extracts_docx_paragraphs_headings_and_tables(self) -> None:
        from docx import Document

        with tempfile.TemporaryDirectory() as directory:
            temp = Path(directory)
            source = temp / "source.docx"
            output = temp / "OFF-001.md"
            document = Document()
            document.add_heading("Course topic", level=1)
            document.add_paragraph("A searchable explanation.")
            table = document.add_table(rows=2, cols=2)
            table.cell(0, 0).text = "Term"
            table.cell(0, 1).text = "Meaning"
            table.cell(1, 0).text = "A | B"
            table.cell(1, 1).text = "First\nSecond"
            document.save(source)

            warnings = extract(source, output)
            text = output.read_text(encoding="utf-8")

            self.assertEqual(warnings, 0)
            self.assertIn("# Extracted source: OFF-001", text)
            self.assertIn("## Course topic", text)
            self.assertIn("A searchable explanation.", text)
            self.assertIn("| Term | Meaning |", text)
            self.assertIn(r"| A \| B | First<br>Second |", text)

    def test_extracts_pptx_slide_text_table_and_notes(self) -> None:
        from pptx import Presentation
        from pptx.util import Inches

        with tempfile.TemporaryDirectory() as directory:
            temp = Path(directory)
            source = temp / "source.pptx"
            output = temp / "OFF-002.md"
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[5])
            slide.shapes.title.text = "Architecture"
            table_shape = slide.shapes.add_table(
                2, 2, Inches(1), Inches(2), Inches(6), Inches(1)
            )
            table_shape.table.cell(0, 0).text = "Layer"
            table_shape.table.cell(0, 1).text = "Role"
            table_shape.table.cell(1, 0).text = "Wiki"
            table_shape.table.cell(1, 1).text = "Knowledge"
            slide.notes_slide.notes_text_frame.text = "Explain the boundary."
            presentation.save(source)

            warnings = extract(source, output)
            text = output.read_text(encoding="utf-8")

            self.assertEqual(warnings, 0)
            self.assertIn("## Slide 1", text)
            self.assertIn("Architecture", text)
            self.assertIn("| Layer | Role |", text)
            self.assertIn("### Speaker notes", text)
            self.assertIn("Explain the boundary.", text)

    def test_marks_picture_slide_for_visual_inspection(self) -> None:
        from base64 import b64decode
        from pptx import Presentation
        from pptx.util import Inches

        one_pixel_png = b64decode(
            "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR4nGNgYAAAAAMA"
            "ASsJTYQAAAAASUVORK5CYII="
        )
        with tempfile.TemporaryDirectory() as directory:
            temp = Path(directory)
            source = temp / "visual.pptx"
            picture = temp / "pixel.png"
            output = temp / "OFF-003.md"
            picture.write_bytes(one_pixel_png)
            presentation = Presentation()
            slide = presentation.slides.add_slide(presentation.slide_layouts[6])
            slide.shapes.add_picture(str(picture), Inches(1), Inches(1))
            presentation.save(source)

            warnings = extract(source, output)

            self.assertEqual(warnings, 1)
            self.assertIn(VISUAL_WARNING, output.read_text(encoding="utf-8"))

    def test_rejects_unsupported_extension(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            temp = Path(directory)
            source = temp / "source.txt"
            source.write_text("text", encoding="utf-8")

            with self.assertRaisesRegex(ValueError, r"\.docx, \.pptx"):
                extract(source, temp / "output.md")


if __name__ == "__main__":
    unittest.main()
